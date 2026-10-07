"""Build + validate compact RE2-safe GSC regex packs for Vametra AI.

Stem-based (not phrase lists) so each pack stays far under GSC's 4,096-char cap
while matching many more real queries. Validated against the 2,149-keyword library.
"""
import re
import openpyxl

PACKS = {
"01 Brand": r"(?i)(vametra|vametr|vamet?ra ?ai|vametraai|vemetra|vametras)",

"02 Duty, tariff & customs": r"(?i)((custom|customs|import|export|basic|bcd|igst|cess|anti.?dumping|safeguard|countervailing|preferential|zero|nil|reduced) ?(duty|duties|tariff|tariffs|rate|rates|calculator|calculation|structure|exemption)|duty (calculator|rate|structure|free|drawback|benefit|exemption|on)|tariff (calculator|rate|code|schedule|classification|by country|on|for|india|usa|uae|eu|uk)|customs (clearance|procedure|process|documentation|valuation|notification|circular|tariff|data|broker|house agent|duty|classification|compliance|documents|records|regulation)|(customs|trade|international trade) compliance|icegate|cbic|hsn? ?duty|how much (duty|customs|tariff)|what is the (duty|tariff|customs duty))",

"03 HS / HSN classification": r"(?i)((hsn?|hs|hts|itc.?hs|commodity|tariff|customs) ?(code|codes|number|classification|chapter|heading|finder|search|lookup|list|directory)|harmoni[sz]ed system|8 ?digit code|6 ?digit code|(find|correct|right|search|check|which|what) .{0,25}(hsn?|hs|hts) ?code|hsn? ?(for|of) )",

"04 Buyers, importers & verification": r"(?i)((buyer|buyers|importer|importers|customer|customers|supplier|suppliers|manufacturer|manufacturers|exporter|exporters|distributor|distributors)s? ?(finder|search|list|lists|leads|data|database|databases|directory|discovery|intelligence|verification|validation|kyc|risk|risk assessment|trust score|background check|due diligence|screening|compliance screening|contacts|enquiry|enquiries|inquiry|inquiries|by country|by product|for export|for india|sourcing from india|outreach)|(find|get|search|source|contact|list of|database of|directory of|verified|genuine|real|authentic|international|foreign|overseas|global|b2b|potential|active|bulk|wholesale|top|indian|government|public sector|trade) ?(buyer|buyers|importer|importers|customer|customers|distributor|distributors|supplier|suppliers|manufacturer|manufacturers|exporter|exporters)|(verify|verification|vet|screen|screening|due diligence|trust score|credit check) .{0,25}(buyer|buyers|importer|importers|supplier|suppliers|company|companies|exporter)|(buyer|importer|supplier|export|import|trade) .{0,20}(scam|scams|fraud|fraudulent|legit|legitimate|trustworthy|safe|genuine)|sanctions (screening|check|list|screened)|denied party|(export|import|trade|sales|business) leads|(export|government procurement|trade) buyers|buyers for (products|export)|b2b marketplace|wholesale marketplace|supplier (directory|discovery|intelligence))",

"05 Landed cost, pricing & Incoterms": r"(?i)(landed cost|landing cost|total cost of import|(export|import|product|quotation|quote|deal|trade) ?(pric(e|ing)|cost|costing|calculator|calculation|margin|margins|profit|profitability|quotation|quote tool)|incoterm|incoterms|inco terms|\b(fob|cif|cfr|exw|fca|fas|cip|cpt|dap|dpu|ddp|dat)\b|freight (cost|rate|rates|calculator|charges|forwarder)|shipping (cost|rate|rates|charges|calculator|line)|cost (sheet|breakup|break.?up) .{0,20}export|how to (price|quote) .{0,20}(export|import|product)|(marine|marine cargo) insurance|cha charges|container (rate|cost|charges|load|volume) ?(calculator)?|cbm calculator|(currency|exchange|fx) (rate|rates|exchange)|exchange rate|usd ?(inr|aed|aud|cny|jpy|sgd)|eur ?usd|gbp ?usd|fx for exporter|dual currency quote|proforma invoice)",

"06 Documents, compliance & schemes": r"(?i)(iec ?(code|certificate|registration|number|apply|online|renewal|modification)?|import export code|importer exporter code|dgft|rcmc|registration cum membership|export promotion (council|bureau)|epc registration|fieo|apeda|spices board|epch|aepc|texprocil|pharmexcil|cdsco|who ?gmp|rodtep|rosctl|meis|seis|duty drawback|advance authori[sz]ation|epcg|moowr|sez unit|interest (equalisation|equalization|subvention)|market access initiative|mai scheme|ties scheme|trade infrastructure for export|export (incentive|incentives|subsidy|subsidies|scheme|schemes)|government export scheme|export (document|documents|documentation|paperwork|formalities|procedure|process|compliance|licence|license|registration|certification|certificate|invoice|declaration|clearance|contract)|import (document|documents|documentation|paperwork|formalities|procedure|process|compliance|licence|license|registration)|documents? (needed|required|for export|for import)|(bank|shipping|international shipping) documents|shipping bill|bill of lading|airway ?bill|commercial invoice|packing list|purchase order|certificate of origin|coo certificate|rules of origin|letter of credit|letter of undertaking|\bl ?c\b|payment terms|\b(da|dp|tt)\b payment|bank realisation|brc|e.?brc|gst ?(registration|refund|lut|export|exports|on export|on exports|for exporter|invoice)|gstin|lut (bond|for export)|phytosanitary|fumigation|halal certificat|bis (certification|registration)|saber|saso|gacc|ce mark|fda registration|cbam|eudr|reach compliance|kyc .{0,15}(buyer|importer)|customs broker|customs house agent|freight forwarder|ecgc|export credit|trade finance|export insurance|(company|llp|private limited|shop act|business) registration|exporter (licence|license|registration))",

"07 FTAs, corridors & countries": r"(?i)(\bfta\b|free trade agreement|preferential trade|cepa|cepta|ceca|ecta|rcep|gsp|trade (agreement|agreements|corridor|corridors|route|lane|relations|partner|partners|deal|profile)|(export|exporting|import|importing|sell|ship|send|source|freight|shipping|buyers?|importers?|tariff|market entry) .{0,30}(to|from|into|for|in)? ?(india|uae|dubai|usa|u\.s\.|united states|america|uk|united kingdom|britain|europe|eu|germany|france|italy|netherlands|spain|poland|china|japan|korea|singapore|malaysia|indonesia|vietnam|thailand|bangladesh|sri lanka|nepal|saudi|ksa|qatar|oman|kuwait|bahrain|egypt|africa|nigeria|kenya|south africa|australia|new zealand|canada|mexico|brazil|russia|turkey|israel|armenia)|(india|indian|uae|dubai|usa|us|uk|armenia|australia|china|singapore|saudi|germany|netherlands|bangladesh) .{0,20}(export|exports|import|imports|import export|trade|tariff|duty|buyers|importers|suppliers|trade profile|market entry)|(country|countries) .{0,20}(buyers|importers|imports|exports|import export|trade profile|selection|profile)|(export|exports|import|imports) by country|(india|usa|uae|uk|china) ?(usa|uae|uk|china)? (freight|shipping)|find export market by country|top (export|import) (market|markets|destination|destinations|partner|partners)|(best|which) (country|countries|market|markets) .{0,25}(export|import|sell|demand)|market entry ?(plan|strategy|consultant|consulting|services|service)?)",

"08 Products & market research": r"(?i)((export|import|demand|market|sourcing|trade|global|product) .{0,25}(potential|research|analysis|trend|trends|demand|opportunity|opportunities|size|data|statistics|stats|volume|selection|insight|insights)|(which|what|best|top|most|profitable|high.?demand|low.?competition) .{0,25}(product|products|commodity|commodities|item|items|market|markets|export market|export markets|export products) ?.{0,25}(export|import|demand|sell|profit)?|product (research|discovery|sourcing|classification|profile|demand|opportunity)|(agriculture|agri|food|food processing|chemical|chemicals|engineering|fmcg|textile|textiles|pharma|pharmaceutical|pharmaceuticals|handicraft|leather|jewellery|jewelry|automotive|auto) .{0,20}(buyer|buyers|importer|importers|export|exports|exporter|exporters|export market|export markets|demand)|basmati|rice export|agarbatti|incense stick|spice|spices|turmeric|chilli|cumin|cotton|garment|apparel|medicine export|tea export|coffee export|sugar export|wheat export|onion export|seafood|shrimp|buffalo meat|engineering goods|auto (parts|component)|organic (food|product)|international sourcing|source products globally|manufacturer (data|database))",

"09 Expos, events & trade fairs": r"(?i)(trade (fair|fairs|show|shows|expo|expos|exhibition|exhibitions|event|events|delegation|mission)|business (expo|expos|exhibition|exhibitions)|import export expo|(expo|exhibition|fair|show) .{0,25}(india|uae|dubai|usa|germany|europe|china|calendar|list|schedule|2026|2027|buyer|buyers|export|import)|gulfood|anutec|canton fair|india international trade fair|iitf|aahar|biofach|gitex|arab health|interpack|ambiente|sial|foodex|buyer seller meet|b2b matchmaking)",

"10 AI, tools, platform & competitors": r"(?i)(\bai\b .{0,30}(trade|export|import|customs|duty|buyer|buyers|supplier|hsn?|hs code|landed cost|compliance|tariff|sourcing|research)|(trade|export|import|customs|buyer|sourcing) .{0,20}\bai\b|ai (tool|tools|platform|software|assistant|agent|copilot) .{0,25}(trade|export|import|business)|chatgpt for (exporter|exporters|importer|importers|global trade|trade)|(best|top|free) .{0,25}(tool|tools|software|platform|app|portal|dashboard|solution|system) .{0,25}(export|import|trade|exporter|importer|customs|buyer)|(trade|import.?export|customs|buyer|market|supply chain|competitive|global trade) (intelligence|data|database|analytics|platform|dashboard|insight|insights|management|management software|management system|operating system|operations platform|risk|analysis)|(export|import|trade|global trade) (tools|software|platform|management system|command center)|(exporter|importer) (platform|data|database)|importer.?exporter database|(shipment|customs shipment|import) records|shipment date data|shipment data|pricing trends|international quote tool|bill of entry data|customs data|global trade data|export import data|trade statistics portal|competitor analysis|trade (ecosystem|assurance)|volza|panjiva|importyeti|import yeti|trademo|zauba|tradeatlas|export ?genius|eximpedia|tendata|seair|infodrive|connect2india|alibaba|indiamart|made.?in.?china|tradeindia|\balternative\b|\bcompetitors?\b|\bvs\b .{0,20}(volza|panjiva|importyeti|trademo|zauba|alibaba|indiamart|eximpedia|tendata))",

"11 Beginner & business setup": r"(?i)((start|starting|how to start|setup|set up|begin|launch|learn) .{0,30}(export|import|export.?import|exim|trading|international) .{0,15}(business|company|firm|agency|trade)?|export business (plan|idea|ideas|from home|for beginners|without investment|profit|licence|license|registration|requirements|startup|process|india)|import.?export business ?(india|registration|requirements|startup|plan|profitable)?|import business (plan|idea|ideas|profit|consultant)|exim (business|course|training|policy)|(export|exports|import|imports|international trade) .{0,20}(for beginners|step by step|beginner|guide|tutorial|course|training|certification|learn|basics|101|meaning)|(first|first time|beginner) (export|exporter|shipment|order)|(small|home|msme|udyam) .{0,20}(export|import|business)|how to become (an )?(exporter|importer)|(export|import) (profit|margin|roi|investment|capital)|(capital|investment) (required )?for export|trade academy|exporter training|(import export|international trade|global trade|export|import)$|export(ing|ers?)? (business|india|businesses abroad)|import(ing|ers?)? business)",

"12 Question-intent (GEO / AI answers)": r"(?i)((how (do|can|to|much|many|long|does)|what (is|are|should|does|documents?|certificat|duty|tariff|hsn?|products?|countr)|which (country|countries|product|products|market|markets|buyer|supplier|hsn?|hs|tool|platform|expo)|where (can|do|to) .{0,15}(find|get|buy|sell|source|verify|check)|who (is|are|can) .{0,25}(buyer|buyers|importer|exporter|supplier)|can (i|we|you) .{0,25}(export|import|sell|ship|source|start|verify|find)|is (there|it|my|this) .{0,30}(tool|platform|ai|legal|safe|legit|possible|required|mandatory|duty.?free)|do i need|should i|why (do|is|are)) .{0,60}(export|import|trade|customs|duty|tariff|hsn?|hs code|buyer|buyers|importer|supplier|document|certificate|landed cost|incoterm|fta|cepa|expo|exhibition|india|overseas|international|market|product))",

"13 Consulting & advisory services": r"(?i)((export|import|import.?export|trade|customs|sourcing|compliance|market entry|international|global trade|buyer) .{0,20}(consultant|consultancy|consulting|consulting services|advisor|advisory|audit|agency|service|services|expert|specialist)|(buyer|supplier|manufacturer|product) (outreach|sourcing|lead generation) ?(service|services)?|lead generation service|trade compliance (audit|consulting)|export (strategy|market) consultant)",

"14 Trade news, policy, prices & freight": r"(?i)((trade|customs|export|import|shipping|freight|tariff|supply chain|global trade|international trade|global supply chain|policy|regulation) .{0,15}(news|update|updates|alert|alerts|policy|regulation|restrictions)|trade (war|sanctions|restrictions|policy)|global tariff|(brent|wti) crude|commodity (price|prices)|(gold|silver|copper|natural gas) price|(ocean|air|sea) freight (trend|trends|rate|rates|data)|red sea (shipping|rerouting|disruption)|port congestion|container (shortage|availability))",
}


def validate():
    wb = openpyxl.load_workbook('/tmp/kw.xlsx', read_only=True)
    lib = [r[1] for r in wb['SEO Keyword Library'].iter_rows(min_row=5, values_only=True) if r and r[1]]
    noise = {str(r[0]).lower() for r in wb['Excluded Noise'].iter_rows(min_row=5, values_only=True) if r and r[0]}
    compiled = {}
    print(f"{'PACK':38s} {'CHARS':>6s}  STATUS")
    for name, rx in PACKS.items():
        assert len(rx) < 4096, f"{name} too long"
        try:
            compiled[name] = re.compile(rx[4:], re.I) if rx.startswith("(?i)") else re.compile(rx, re.I)
            status = "ok"
        except re.error as e:
            status = f"REGEX ERROR {e}"
        bad = [t for t in ("(?=", "(?!", "(?<", r"\1", r"\2") if t in rx]
        if bad:
            status += f" RE2-UNSAFE {bad}"
        print(f"{name:38s} {len(rx):6d}  {status}")

    matched, unmatched = 0, []
    hits = {k: 0 for k in PACKS}
    for kw in lib:
        k = str(kw).strip()
        hit = [n for n, c in compiled.items() if c.search(k)]
        if hit:
            matched += 1
            for n in hit:
                hits[n] += 1
        else:
            unmatched.append(k)
    print(f"\nLIBRARY COVERAGE: {matched}/{len(lib)} = {matched / len(lib) * 100:.1f}%")
    print("per-pack hits:", {n: v for n, v in sorted(hits.items())})
    real_unmatched = [u for u in unmatched if u.lower() not in noise]
    print(f"\nUNMATCHED (excluding the 50 noise terms): {len(real_unmatched)}")
    for u in real_unmatched[:60]:
        print("  -", u)


if __name__ == "__main__":
    validate()
